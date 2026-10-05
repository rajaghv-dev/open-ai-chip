"""Approximate text-fit checker (no renderer needed): wraps each text box with Arial metrics and
compares the needed height with the box height. Usage: python check_fit.py deck.pptx"""
import sys
from pptx import Presentation
from pptx.util import Emu
from PIL import ImageFont
D = "/System/Library/Fonts/Supplemental/"
cache = {}
def font(sz, bold):
    k = (round(sz * 4), bold)
    if k not in cache: cache[k] = ImageFont.truetype(D + ("Arial Bold.ttf" if bold else "Arial.ttf"), int(round(sz * 10)))
    return cache[k]
def width(t, sz, bold): return font(sz, bold).getlength(t) / 10.0
def lines_needed(par_runs, wpt):
    # par_runs: list of (text, size, bold); greedy word wrap
    words = []
    for t, sz, b in par_runs:
        for i, w in enumerate(t.replace("\n", " \n ").split(" ")):
            words.append((w, sz, b))
    n, cur, maxsz = 1, 0.0, max([r[1] for r in par_runs] or [12])
    for w, sz, b in words:
        if w == "\n": n += 1; cur = 0; continue
        ww = width(w + " ", sz, b)
        if cur + ww - width(" ", sz, b) > wpt and cur > 0: n += 1; cur = ww
        else: cur += ww
    return n, maxsz
def need(tf, wpt):
    h = 0
    for p in tf.paragraphs:
        runs = [(r.text, (r.font.size.pt if r.font.size else 12), bool(r.font.bold)) for r in p.runs]
        if not runs: h += 12 * 1.2; continue
        # explicit line breaks inside runs
        n, sz = lines_needed(runs, wpt); h += n * sz * 1.2
    return h
prs = Presentation(sys.argv[1]); bad = 0
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            tf = sh.text_frame
            ml = (tf.margin_left or 0) / 12700; mr = (tf.margin_right or 0) / 12700; mt = (tf.margin_top or 0) / 12700; mb = (tf.margin_bottom or 0) / 12700
            w = sh.width / 12700 - ml - mr; h = sh.height / 12700 - mt - mb
            nh = need(tf, w)
            if nh > h * 1.02:
                bad += 1; print(f"slide {i}: OVER {nh:.0f}pt > {h:.0f}pt  [{sh.name}] {tf.text[:60]!r}")
        if sh.has_table:
            t = sh.table; tot = 0
            for r in t.rows:
                rh = 0
                for ci, c in enumerate(r.cells):
                    w = t.columns[ci].width / 12700 - (c.margin_left + c.margin_right) / 12700
                    rh = max(rh, need(c.text_frame, w) + (c.margin_top + c.margin_bottom) / 12700)
                tot += max(rh, r.height / 12700)
            bottom = sh.top / 12700 + tot
            print(f"slide {i}: table needs {tot/72:.2f} in high, bottom at {bottom/72:.2f} in (slide 5.625, footnote/footer zone starts ~4.7)")
print("overflow boxes:", bad)
